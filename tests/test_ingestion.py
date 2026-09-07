from idiolect.ingestion import clean_text, ingest
from idiolect.models import Document


def test_clean_text():
    dirty = "This  is\t\ta   test\n\n\n\nwith   spaces."
    cleaned = clean_text(dirty)
    assert "  " not in cleaned
    assert "\t" not in cleaned
    assert cleaned == "This is a test\n\nwith spaces."


def test_ingest():
    sample = "The quick brown fox jumps over the lazy dog. It was a very agile fox!"
    doc = ingest(sample)
    assert isinstance(doc, Document)
    assert doc.word_count > 0
    assert doc.sentence_count == 2
    assert len(doc.tokens) > 0
    assert doc.spacy_doc is not None


def test_find_text_files_and_ingest_path(tmp_path):
    from idiolect.ingestion import find_text_files, ingest_path

    # Directory with multiple supported text files and ignored files
    dir_path = tmp_path / "corpus"
    dir_path.mkdir()

    f1 = dir_path / "chapter1.txt"
    f1.write_text("The journey began early in the crisp morning air.")

    f2 = dir_path / "chapter2.md"
    f2.write_text("Across the wide valley, shadows grew as night approached.")

    hidden = dir_path / ".hidden.txt"
    hidden.write_text("Should be ignored.")

    ignored = dir_path / "data.csv"
    ignored.write_text("col1,col2\n1,2")

    found = find_text_files(dir_path)
    assert len(found) == 2
    assert found[0].name == "chapter1.txt"
    assert found[1].name == "chapter2.md"

    # Single file
    assert find_text_files(f1) == [f1]

    # Non-existent
    assert find_text_files(tmp_path / "does_not_exist") == []

    # Ingest path on directory
    combined_doc = ingest_path(dir_path)
    assert "The journey began early" in combined_doc.cleaned_text
    assert "Across the wide valley" in combined_doc.cleaned_text
    assert combined_doc.sentence_count >= 2

    # Ingest path on single file
    single_doc = ingest_path(f1)
    assert "The journey began early" in single_doc.cleaned_text


def test_docx_ingestion(tmp_path):
    import docx

    from idiolect.ingestion import extract_text_from_docx, find_text_files, ingest_file

    docx_path = tmp_path / "submission.docx"
    doc = docx.Document()
    doc.add_paragraph("Stylometry provides deep insight into authentic voice.")
    doc.add_paragraph("Consistent syntax patterns distinguish distinct authors.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cell A"
    table.rows[0].cells[1].text = "Cell B"
    doc.save(docx_path)

    extracted = extract_text_from_docx(docx_path)
    assert "Stylometry provides deep insight" in extracted
    assert "Consistent syntax patterns" in extracted
    assert "Cell A | Cell B" in extracted

    # Verify discoverable by find_text_files
    files = find_text_files(tmp_path)
    assert docx_path in files

    # Verify ingest_file
    document = ingest_file(docx_path)
    assert document.word_count > 0
    assert "Stylometry provides deep insight" in document.cleaned_text


def test_pdf_ingestion(tmp_path):
    from fpdf import FPDF

    from idiolect.ingestion import extract_text_from_pdf, find_text_files, ingest_file

    pdf_path = tmp_path / "essay.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="This is an authentic student essay submitted in PDF format.")
    pdf.output(str(pdf_path))

    extracted = extract_text_from_pdf(pdf_path)
    assert "authentic student essay" in extracted

    # Verify discoverable by find_text_files
    files = find_text_files(tmp_path)
    assert pdf_path in files

    # Verify ingest_file
    document = ingest_file(pdf_path)
    assert document.word_count > 0
    assert "authentic student essay" in document.cleaned_text


def test_invalid_and_empty_extraction(tmp_path):
    import pytest

    from idiolect.ingestion import extract_text_from_docx, extract_text_from_pdf

    bad_docx = tmp_path / "corrupt.docx"
    bad_docx.write_bytes(b"not a valid zip file")
    with pytest.raises(ValueError, match="Failed to extract text from Word document"):
        extract_text_from_docx(bad_docx)

    bad_pdf = tmp_path / "corrupt.pdf"
    bad_pdf.write_bytes(b"not a valid pdf")
    with pytest.raises(ValueError, match="Failed to extract text from PDF document"):
        extract_text_from_pdf(bad_pdf)
